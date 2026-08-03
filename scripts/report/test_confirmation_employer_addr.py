#!/usr/bin/env python3
"""确认函：归属公司非申请人时地址取值的复现/回归测试。

复现 bug（案例 20688-0018US00）：发明人归属公司用半角括号 "BMC(TIANJIN)MEDICAL CO., LTD."，
而 assignments 转让方是全角括号 "BMC（TIANJIN）MEDICAL CO., LTD."，旧 _norm 仅 strip+lower
匹配不上 → 该组发明人地址取不到（只剩劳动合同归属）。修后：公司名归一去非字母数字，全/半角括号
视同 → 从转让方 company_address 取到地址。

  cd scripts/report && python3 test_confirmation_employer_addr.py
"""
import os
import sys
import tempfile

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)

from submission.confirmation_sheet import generate_confirmation_sheet
from docx import Document

_passed = 0


def ok(cond, msg):
    global _passed
    assert cond, f"FAIL: {msg}"
    _passed += 1
    print(f"  ok - {msg}")


FIELDS = {
    "case_volume": "20688-0018US00", "application_type": "发明", "ae_methods": "None",
    "case_info": {"case_name_en": "RESPIRATORY TREATMENT"},
    "applicants": [{"applicant_name_en": "BMC MEDICAL CO., LTD.", "address1": "Room 110",
                    "city_en": "Beijing", "postcode": "100142", "country_en_us": "China", "entity_en": "Large"}],
    "inventors": [
        {"last_name_en": "WANG", "first_name_en": "Q", "full_name_en": "WANG, Q", "employer": "BMC MEDICAL CO., LTD."},
        {"last_name_en": "CHEN", "first_name_en": "Xi", "full_name_en": "CHEN, Xi", "employer": "BMC(TIANJIN)MEDICAL CO., LTD."},
    ],
    "priorities": [], "parent_applications": [], "filing_route": "Bypass Continuation",
    "assignments": [
        {"assignment_type": "个人转公司", "assignor_inventors": ["WANG, Q", "CHEN, Xi"], "assignee": "BMC MEDICAL CO., LTD."},
        # 转让方用全角括号，与发明人 employer 的半角括号不同（bug 触发点）
        {"assignment_type": "公司转公司", "assigning_company": "BMC（TIANJIN）MEDICAL CO., LTD.",
         "company_representative": "XU", "representative_title": "Mgr",
         "company_address": "3/F, Building No.4, No.1 Xinxing Road, Wuqing District, Tianjin, 301799, China"},
    ],
}


def _doc_text(path):
    d = Document(path)
    parts = [p.text for p in d.paragraphs]
    for t in d.tables:
        for row in t.rows:
            parts.append(" | ".join(c.text for c in row.cells))
    return "\n".join(parts)


def test_non_applicant_employer_address():
    out = tempfile.mkdtemp(prefix="conf_")
    r = generate_confirmation_sheet(FIELDS, out)
    txt = _doc_text(r["file_path"])
    ok("Xinxing Road" in txt, "归属公司(全/半角括号不同)仍取到转让方地址：天津地址出现在确认函")
    ok("CHEN" in txt, "发明人 CHEN 出现")


if __name__ == "__main__":
    test_non_applicant_employer_address()
    print(f"\nAll {_passed} checks passed.")
