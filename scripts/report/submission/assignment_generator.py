#!/usr/bin/env python3
"""Assignment docx 生成（消费 fields_json，移植自 old_submission/assignment/generate_assign.py）。

与旧实现的差异：
- 输入直接用 submission_tasks.fields_json（单一数据源），不再 docx 反解析 / DeepSeek
- 按 assignments 数组逐份生成（一段转让一份文件，支持转让链多段）
- 模板分区删除按页标记定位（个人转公司转让人页 / 公司转公司转让人页 / 公司转公司受让人页），
  不再用硬编码段落索引
- 20618 客户按转让类型选择专用模板；签字区与 Mailing Address 沿用通用模板版式和填充规则
"""
import copy
import re
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import _Cell, _Row

from .common import (
    ASSIGNMENT_TEMPLATE,
    blank_none,
    company_addr_dict,
    _assignee_from_applicant,
    application_type_words,
    build_priority_rows,
    effective_assignments,
    get_assignee,
    inventor_names,
    is_special_client,
    resolve_assignee,
    split_address,
)

CHINESE_PATTERN = re.compile(r"[一-鿿]")

MARKER_PERSONAL = "个人转公司转让人页"
MARKER_COMPANY = "公司转公司转让人页"
MARKER_ASSIGNEE = "公司转公司受让人页"

SPECIAL_ASSIGNMENT_DIR = ASSIGNMENT_TEMPLATE.parent / "20618"
SPECIAL_ASSIGNMENT_TEMPLATES = {
    "个人转公司": SPECIAL_ASSIGNMENT_DIR / "Assignment-20618-inventors.docx",
    "公司转公司": SPECIAL_ASSIGNMENT_DIR / "Assignment-20618-company.docx",
}


# ─── 文本替换与清理（行为对齐旧 generate_assign.py） ────────

def _remove_highlighting(paragraph):
    for run in paragraph.runs:
        if run._element.rPr is not None:
            highlight = run._element.rPr.find(qn("w:highlight"))
            if highlight is not None:
                run._element.rPr.remove(highlight)


def _replace_in_paragraph(paragraph, old, new):
    if old in paragraph.text:
        _remove_highlighting(paragraph)
        paragraph.text = paragraph.text.replace(old, str(new))


def _replace_everywhere(doc, old, new):
    for paragraph in doc.paragraphs:
        _replace_in_paragraph(paragraph, old, new)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _replace_in_paragraph(paragraph, old, new)
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            _replace_in_paragraph(paragraph, old, new)


def _remove_chinese_tables(doc):
    """删除仍含中文的表格（未填充的占位表：多余发明人表、未走分支的公司表等）。"""
    for table in list(doc.tables):
        text = "".join(cell.text for row in table.rows for cell in row.cells)
        if CHINESE_PATTERN.search(text):
            table._element.getparent().remove(table._element)


def _remove_chinese_paragraphs(doc):
    """删除仍含中文的段落（页标记、未替换的占位符行）。"""
    for paragraph in list(doc.paragraphs):
        if CHINESE_PATTERN.search(paragraph.text):
            paragraph._element.getparent().remove(paragraph._element)


def _cleanup_blank_paragraphs(doc):
    """清理删表后残留的空白：

    1. 删除空段落里的显式分页符（删表后分页符会撑出空白页）
    2. 连续空段落收敛为 1 个（保留模板原有的单个间隔段落）
    3. 文末（sectPr 前）的空段落全部删除
    """
    body = doc.element.body
    w_p, w_t, w_br, w_type = qn("w:p"), qn("w:t"), qn("w:br"), qn("w:type")

    def is_empty_p(el):
        return el.tag == w_p and not "".join(t.text or "" for t in el.findall(".//" + w_t)).strip()

    # 1. 空段落里的分页符
    for el in list(body):
        if is_empty_p(el):
            for br in el.findall(".//" + w_br):
                if br.get(w_type) == "page":
                    br.getparent().remove(br)

    # 2. 连续空段落收敛为 1 个
    prev_empty = False
    for el in list(body):
        if is_empty_p(el):
            if prev_empty:
                body.remove(el)
            else:
                prev_empty = True
        elif el.tag == w_p or el.tag.endswith("}tbl"):
            prev_empty = False

    # 3. 文末空段落全删（倒序到第一个有内容的元素为止）
    for el in reversed(list(body)):
        if el.tag.endswith("}sectPr") or el.tag.endswith("}bookmarkEnd"):
            continue
        if is_empty_p(el):
            body.remove(el)
        else:
            break


def _strip_all_highlights(doc):
    for paragraph in doc.paragraphs:
        _remove_highlighting(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    _remove_highlighting(paragraph)


# ─── 按页标记的分区删除 ──────────────────────────────────

def _body_marker_positions(doc):
    """返回 body 子元素列表及三个页标记的下标。"""
    body = doc.element.body
    children = list(body)
    positions = {}
    for idx, el in enumerate(children):
        if el.tag == qn("w:p"):
            text = "".join(t.text or "" for t in el.findall(".//" + qn("w:t"))).strip()
            if text in (MARKER_PERSONAL, MARKER_COMPANY, MARKER_ASSIGNEE):
                positions[text] = idx
    return body, children, positions


def _remove_sections(doc, assignment_type, special_client):
    """按转让类型保留对应签字分区，整段删除（含段落和表格）其余分区。

    模板布局：正文 → [M1 个人转公司转让人页 → 发明人签字表×8]
                 → [M2 公司转公司转让人页 → 转让公司表]
                 → [M3 公司转公司受让人页 → 受让公司表]
    """
    body, children, pos = _body_marker_positions(doc)
    m1, m2, m3 = pos.get(MARKER_PERSONAL), pos.get(MARKER_COMPANY), pos.get(MARKER_ASSIGNEE)
    if m1 is None or m2 is None or m3 is None:
        raise RuntimeError(f"Assignment 模板缺少页标记：{pos.keys()}，模板可能被改动")

    if assignment_type == "个人转公司":
        # 保留 M1..M2 发明人区；删除 M2 起全部（公司转让 + 受让）
        to_remove = children[m2:]
    elif special_client:
        # 公司转公司 + 特例客户：删除个人区，保留转让公司 + 受让公司
        to_remove = children[m1:m2]
    else:
        # 公司转公司：删除个人区 + 受让区
        to_remove = children[m1:m2] + children[m3:]

    for el in to_remove:
        if el.tag == qn("w:sectPr"):  # 节属性保留，否则页面设置丢失
            continue
        body.remove(el)


# ─── 分区填充 ────────────────────────────────────────────

def _fill_priority_table(doc, rows):
    for table in doc.tables:
        header_text = "".join(cell.text for cell in table.rows[0].cells)
        if "Application Number" in header_text or "Country of Filing" in header_text:
            while len(table.rows) > 1:
                table._tbl.remove(table.rows[1]._tr)
            for app_no, country, date in rows:
                new_row = table.add_row()
                if len(new_row.cells) >= 3:
                    new_row.cells[0].text = app_no
                    new_row.cells[1].text = country
                    new_row.cells[2].text = date
            return


def _clone_address_rows(table, ref_row, extra_lines):
    """克隆地址行（连同底边线/格式）插在其后，逐行填入多出的地址行。"""
    prev_tr = ref_row._tr
    for line in extra_lines:
        new_tr = copy.deepcopy(ref_row._tr)
        prev_tr.addnext(new_tr)
        prev_tr = new_tr
        new_row = _Row(new_tr, table)
        # 地址在最后一格：替换文字（底边线在 tcPr，cell.text 不影响）
        addr_cell = new_row.cells[-1]
        addr_cell.text = line
        # 文字落在底边线上（对齐地址1格的 vAlign=bottom，否则贴顶边）
        addr_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM


def _inventor_signature_tables(doc):
    """所有含"发明人打印版名字"占位的签字表（含克隆出来的）。"""
    return [
        table for table in doc.tables
        if any("发明人打印版名字" in cell.text for row in table.rows for cell in row.cells)
    ]


def _preceding_spacer_paragraphs(tbl_el):
    """签字表前紧邻的空行段落（作为表间距样板），按文档顺序返回。"""
    out = []
    prev = tbl_el.getprevious()
    while prev is not None and prev.tag == qn("w:p"):
        out.append(prev)
        prev = prev.getprevious()
    return list(reversed(out))


def _set_page_break_before_table(table):
    """让该签字表从新的一页开始：pageBreakBefore 设在表首格首段（位于表格内，
    不受正文空段清理影响；LibreOffice/Word 均据此另起一页）。幂等。"""
    p = table.rows[0].cells[0].paragraphs[0]
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:pageBreakBefore')) is None:
        pPr.insert(0, OxmlElement('w:pageBreakBefore'))


def _fill_inventor_tables(doc, names, addr_lines):
    """个人转公司：按序把发明人填入签字表（含"发明人打印版名字"占位的表）。

    发明人数超过模板签字表数时，克隆最后一个签字表（连同其前的空行间隔）补足，
    避免超出部分无处签字（A 方案）。
    addr_lines 为地址多行列表：地址1/地址2 填前两行，超 2 行则克隆地址2行补足。
    返回 (填入发明人数, 克隆补足的签字表数)。
    """
    line1 = addr_lines[0] if addr_lines else ""
    line2 = addr_lines[1] if len(addr_lines) > 1 else ""
    extra = addr_lines[2:] if len(addr_lines) > 2 else []

    inventor_tables = _inventor_signature_tables(doc)
    added = 0
    if inventor_tables and len(names) > len(inventor_tables):
        last_tbl = inventor_tables[-1]._tbl
        spacer_tmpls = _preceding_spacer_paragraphs(last_tbl)
        anchor = last_tbl
        for _ in range(len(names) - len(inventor_tables)):
            for sp in spacer_tmpls:  # 先补间距空行，保持表与表间距一致
                new_sp = copy.deepcopy(sp)
                anchor.addnext(new_sp)
                anchor = new_sp
            new_tbl = copy.deepcopy(last_tbl)  # 克隆的是未填充的占位表
            anchor.addnext(new_tbl)
            anchor = new_tbl
            added += 1
        inventor_tables = _inventor_signature_tables(doc)  # 重新收集（含克隆）

    filled = 0
    for name, table in zip(names, inventor_tables):
        addr2_row = None
        for row in table.rows:
            for cell in row.cells:
                if "发明人打印版名字" in cell.text:
                    cell.text = cell.text.replace("发明人打印版名字", name)
                elif "发明人签字日期" in cell.text:
                    cell.text = cell.text.replace("发明人签字日期", "")
                elif "发明人地址1" in cell.text:
                    cell.text = cell.text.replace("发明人地址1", line1)
                elif "发明人地址2" in cell.text:
                    cell.text = cell.text.replace("发明人地址2", line2)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM  # 与地址1一致贴底边
                    addr2_row = row
        if extra and addr2_row is not None:
            _clone_address_rows(table, addr2_row, extra)
        filled += 1
    # 每个发明人签字表各占一页：给第 2 个起的已填签字表设 pageBreakBefore（分页显示）
    used = list(inventor_tables)[:len(names)]
    for _t in used[1:]:
        _set_page_break_before_table(_t)
    return filled, added


# ─── 主入口 ──────────────────────────────────────────────

def generate_assignment(fields: dict, assignment: dict, output_dir, seq: int = 1,
                        total: int = 1, name_tag: str = "") -> dict:
    """按一段转让安排生成一份 Assignment docx。

    :param fields: submission_tasks.fields_json
    :param assignment: assignments 数组中的一个元素
    :param output_dir: 输出目录
    :param seq/total: 当前序号 / 总份数（total>1 时文件名带序号和标签）
    :param name_tag: 文件名标签（多份区分用；如发明人姓）。缺省用 assignment_type。
    :return: {"file_path", "file_name", "warnings": [...]}
    """
    fields = blank_none(fields or {})          # 深度 None→''，Assignment 不出现字面 'None'
    assignment = blank_none(assignment or {})
    assignment_type = assignment.get("assignment_type") or ""
    if assignment_type not in ("个人转公司", "公司转公司"):
        raise ValueError(f"assignment_type 非法：{assignment_type!r}")
    special = is_special_client(fields.get("case_volume") or "")
    template = SPECIAL_ASSIGNMENT_TEMPLATES[assignment_type] if special else ASSIGNMENT_TEMPLATE
    if not template.exists():
        raise RuntimeError(f"缺少模板 {template}")

    warnings = []
    case_volume = fields.get("case_volume") or ""
    title = ((fields.get("case_info") or {}).get("case_name_en") or "").upper()  # 申请名称全部大写
    # 受让人按本段 assignee 名字解析（匹配不到回退首位申请人并告警）
    seg_assignee_name = (assignment.get("assignee") or "").strip()
    assignee = resolve_assignee(fields, seg_assignee_name)
    if seg_assignee_name and not assignee.get("matched"):
        warnings.append(f"受让人「{seg_assignee_name}」未匹配到申请人，已回退首位申请人「{assignee['name']}」")
    # 发明人按本段 assignor_inventors 子集（缺省时兜底全体，兼容存量数据）
    seg_inventors = [str(n).strip() for n in (assignment.get("assignor_inventors") or []) if str(n).strip()]
    names = seg_inventors or inventor_names(fields)
    app_type_eng, confirm_type = application_type_words(fields.get("application_type") or "")
    today = datetime.now()
    today_dot = today.strftime("%Y.%m.%d")

    doc = Document(str(template))

    # 1. 页眉案号
    for section in doc.sections:
        for paragraph in section.header.paragraphs:
            if "案号" in paragraph.text:
                paragraph.text = paragraph.text.replace("案号", case_volume)

    # 2. 让与人 + am/is/are（仅这两处按人数/法人动态；其余斜杠形式保持不动，按用户要求）：
    #    个人转公司 → 1 名 I/am、多名 WE/are；公司转公司 → 公司名 / is（单数法人）
    if assignment_type == "个人转公司":
        _single = len(names) == 1
        _replace_everywhere(doc, "让与人", "I" if _single else "WE")
        _replace_everywhere(doc, "am/is/are", "am" if _single else "are")
    else:
        assigning_company = assignment.get("assigning_company") or ""
        if not assigning_company:
            warnings.append("公司转公司缺少转让方公司名称（assigning_company）")
        _replace_everywhere(doc, "让与人", assigning_company)
        _replace_everywhere(doc, "am/is/are", "is")

    # 3. 发明名称（标题样式）
    for paragraph in doc.paragraphs:
        if "发明名称" in paragraph.text:
            paragraph.text = title
            try:
                paragraph.style = doc.styles["Heading 2"]
            except KeyError:
                pass

    # 4. 受让人名称 / 申请类型措辞
    _replace_everywhere(doc, "个人或公司", assignee["name"])
    _replace_everywhere(doc, "确认类型", confirm_type)
    _replace_everywhere(doc, "申请类型", app_type_eng)
    if special:
        _replace_everywhere(doc, "{{ASSIGNEE_ADDRESS}}", " ".join(assignee["address_lines"]))

    # 5. 优先权表（优先权 + PCT + 美国母案链）
    _fill_priority_table(doc, build_priority_rows(fields))

    # 6. 签字分区填充
    if assignment_type == "个人转公司":
        filled, added = _fill_inventor_tables(doc, names, assignee["address_lines"])
        if added:
            warnings.append(f"发明人 {len(names)} 名，签字表已自动扩容补足 {added} 个")
        if filled < len(names):
            warnings.append(f"发明人 {len(names)} 名超出模板签字表容量且扩容失败，仅填入 {filled} 名")
    else:
        # 公司地址：结构化 company_* 六字段拼全址拆 2 行（同申请人）；缺省回退旧单串
        company_addr = _assignee_from_applicant(company_addr_dict(assignment))
        if special:
            _replace_everywhere(doc, "{{ASSIGNOR_ADDRESS}}", " ".join(company_addr["address_lines"]))
        _replace_everywhere(doc, "转让公司名称", assignment.get("assigning_company") or "")
        _replace_everywhere(doc, "转让公司代表打印版名字", assignment.get("company_representative") or "")
        _replace_everywhere(doc, "转让公司代表职位", assignment.get("representative_title") or "")
        _replace_everywhere(doc, "转让公司地址1", company_addr["address1"])
        _replace_everywhere(doc, "转让公司地址2", company_addr["address2"])
        _replace_everywhere(doc, "转让公司代表签字日期", today_dot)
        if special:
            # 特例客户保留受让人签字页；代表姓名/职位留空待人工补签（不填占位符）
            _replace_everywhere(doc, "受让公司名称", assignee["name"])
            _replace_everywhere(doc, "受让公司代表打印版名字", "")
            _replace_everywhere(doc, "受让公司代表职位", "")
            _replace_everywhere(doc, "受让公司地址1", assignee["address1"])
            _replace_everywhere(doc, "受让公司地址2", assignee["address2"])
            _replace_everywhere(doc, "受让公司代表签字日期", today_dot)

    # 7. 按页标记删除不适用的签字分区（段落+表格整段移除）
    _remove_sections(doc, assignment_type, special)

    # 8. 收尾清理：高亮、残留中文表格/段落（未填充的占位表、页标记）、删表后的空白段/分页符
    _strip_all_highlights(doc)
    _remove_chinese_tables(doc)
    _remove_chinese_paragraphs(doc)
    _cleanup_blank_paragraphs(doc)
    if special:
        unresolved = "\n".join([p.text for p in doc.paragraphs] +
                               [c.text for t in doc.tables for row in t.rows for c in row.cells])
        if re.search(r"\{\{[^}]+\}\}", unresolved):
            raise ValueError("20618 Assignment 模板仍有未填充字段")

    # 9. 保存
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    date_str = today.strftime("%Y-%m-%d")
    if total > 1:
        file_name = f"{date_str}-Assignment-{case_volume}-{seq}-{name_tag or assignment_type}.docx"
    else:
        file_name = f"{date_str}-Assignment-{case_volume}.docx"
    file_path = output_dir / file_name
    doc.save(str(file_path))

    return {"file_path": str(file_path), "file_name": file_name, "warnings": warnings}


def generate_assignments(fields: dict, output_dir) -> list:
    """按转让段逐段生成，返回结果列表。

    转让段取 effective_assignments：发明人有 employer 时按归属公司分组派生个人转公司段
    （与确认函同源，解决"多公司只出一份"），无 employer 时沿用上游 assignments。
    一段转让一份文件：个人转公司段的所有发明人在同一份里、每个发明人签字表**各占一页**
    （分页显示，见 _fill_inventor_tables）；公司转公司段一份。派生告警挂到首份结果 warnings。

    生成前先清掉 output_dir 里旧的 Assignment 产物（docx/pdf），避免旧份（如曾按发明人
    拆成的多份）与新份混显。只清 Assignment，不碰同目录 POA/Transmittal/确认函/Declaration。"""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    for _old in output_dir.glob("*Assignment*"):
        try:
            _old.unlink()
        except OSError:
            pass
    assignments, derive_warnings = effective_assignments(fields)
    if not assignments:
        raise ValueError("assignments 为空，无法生成 Assignment")
    total = len(assignments)
    results = [
        generate_assignment(fields, item, output_dir, seq=i, total=total)
        for i, item in enumerate(assignments, start=1)
    ]
    if derive_warnings and results:
        results[0]["warnings"] = derive_warnings + (results[0].get("warnings") or [])
    return results
