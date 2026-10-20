import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from docx.oxml.ns import qn

from submission.assignment_generator import (SPECIAL_ASSIGNMENT_TEMPLATES,
                                             _special_chain, generate_assignment,
                                             generate_assignments)
from submission.common import ASSIGNMENT_TEMPLATE, full_mailing_address_en


def doc_text(path):
    doc = Document(path)
    text = "\n".join([p.text for p in doc.paragraphs]
                     + [c.text for table in doc.tables for row in table.rows for c in row.cells]
                     + [p.text for section in doc.sections for p in section.header.paragraphs])
    return doc, text


class SpecialAssignmentTemplateTests(unittest.TestCase):
    def test_confirmation_page_mother_case_numbers_enter_chain(self):
        # 确认函页只提供 Application No / Filing Date，不提供 Prior 字段。
        fields = {
            "case_volume": "20618-9999US00",
            "case_info": {"case_name_en": "A NEW DEVICE"},
            "applicants": [{"applicant_name_en": "TEST CORP.", "address1": "1 Main Street"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice"}],
            "filing_route": "Continuation in part of",
            "parent_applications": [
                {"Application No": "18/340,880", "Filing Date": "2023-06-25",
                 "Prior Application Number": ""},
                {"Application No": "17/452,087", "Filing Date": "2021-10-25",
                 "Prior Application Number": "", "Patent Number": "11,686,867"},
            ],
        }
        chain = _special_chain(fields)
        self.assertIn("a continuation-in-part of U.S. application filed on June 25, 2023 "
                      "(U.S. Application No. 18/340,880)", chain)
        self.assertIn("U.S. Patent Application No. 17/452,087, filed on October 25, 2021 "
                      "(now Patent No. 11,686,867)", chain)
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                  "assignee": "TEST CORP."}, folder)
            _, generated_text = doc_text(result["file_path"])
            self.assertIn(chain, generated_text)

    def test_new_chain_does_not_use_this_app_as_missing_prior(self):
        chain = _special_chain({"parent_applications": [
            {"Application No": "18/111,111", "Prior Application Number": "17/222,222"},
            {"Application No": "18/333,333", "Prior Application Number": ""},
        ]})
        self.assertIn("17/222,222", chain)
        self.assertNotIn("18/333,333", chain)

    def test_municipality_is_not_repeated_in_mailing_address(self):
        address = {"address1": "No. 1 Main Road", "city_en": "Shanghai",
                   "province_city_en_us": "shanghai", "postcode": "200000",
                   "country_en_us": "China"}
        self.assertEqual(full_mailing_address_en(address),
                         "No. 1 Main Road, Shanghai, 200000, China")
        address["province_city_en_us"] = "Jiangsu"
        self.assertIn("Shanghai, Jiangsu", full_mailing_address_en(address))

    def test_priority_only_and_pct_not_duplicated(self):
        chain = _special_chain({
            "case_advance_info": {"pct_app_no": "PCT/CN2024/123456",
                                  "pct_app_date": "2024-03-22"},
            "priorities": [
                {"priority_no": "PCT/CN2024/123456", "priority_country_id": "WO"},
                {"priority_no": "CN 202410389713.5", "priority_country_id": "CN",
                 "priority_date": "2024-04-02"},
            ],
        })
        self.assertEqual(chain.count("PCT/CN2024/123456"), 1)
        self.assertIn("Chinese Patent Application No. 202410389713.5 filed on April 2, 2024", chain)
        priority_only = _special_chain({"priorities": [{"priority_no": "202511280001.X",
                                                        "priority_country_id": "CN",
                                                        "priority_date": "2025-11-28"}]})
        self.assertIn("claiming priority to Chinese Patent Application No. 202511280001.X", priority_only)

    def test_us_pct_and_two_cn_priorities_use_customer_prose(self):
        chain = _special_chain({
            "filing_route": "Continuation in part of",
            "parent_applications": [{"Prior Application Number": "18/049,277",
                                     "Filing Date": "2022-10-24",
                                     "Continuity Type": "Continuation in part of"}],
            "case_advance_info": {"pct_app_no": "PCT/CN2021/089454",
                                  "pct_app_date": "2021-04-23"},
            "priorities": [
                {"priority_no": "202010327417.4", "priority_country_id": "CN",
                 "priority_date": "2020-04-23"},
                {"priority_no": "202010374378.3", "priority_country_id": "CN",
                 "priority_date": "2020-05-06"},
            ],
        })
        self.assertEqual(chain,
            "for which is a continuation-in-part of U.S. application filed on October 24, 2022 "
            "(U.S. Application No. 18/049,277), which is a continuation of International Patent "
            "Application No. PCT/CN2021/089454, filed on April 23, 2021; "
            "which claims priority to Chinese Patent Application No. 202010327417.4 "
            "filed on April 23, 2020, and claims priority to Chinese Patent Application No. "
            "202010374378.3 filed on May 6, 2020; and")

    def test_pct_only_routes_and_us_priority(self):
        fields = {
            "case_advance_info": {"pct_app_no": "PCT/CN2024/123456",
                                  "pct_app_date": "2024-03-22"},
            "priorities": [{"priority_no": "18/123,456", "priority_country_id": "US",
                            "priority_date": "2023-03-22"}],
        }
        for route, relation in (("Bypass Continuation", "a Continuation of International application"),
                                ("371 National Stage", "a 371 national-stage entry of International Patent Application")):
            with self.subTest(route=route):
                fields["filing_route"] = route
                chain = _special_chain(fields)
                self.assertIn(f"for which is {relation}", chain)
                self.assertIn("PCT/CN2024/123456", chain)
                self.assertIn("U.S. Patent Application No. 18123456 filed on March 22, 2023", chain)

    def test_customer_five_filing_routes(self):
        cn = lambda no, day: {"priority_no": no, "priority_country_id": "CN", "priority_date": day}
        pct = {"pct_app_no": "PCT/CN2023/081289", "pct_app_date": "2023-03-14"}
        priority = cn("202211097238.1", "2022-09-08")

        con = _special_chain({"filing_route": "Bypass Continuation", "case_advance_info": pct,
                              "priorities": [priority]})
        self.assertEqual(con,
            "for which is a Continuation of International application filed on March 14, 2023 "
            "(International Application No. PCT/CN2023/081289), claiming priority to Chinese "
            "Patent Application No. 202211097238.1 filed on September 8, 2022; and")
        self.assertEqual(_special_chain({"filing_route": "Bypass Continuation", "case_advance_info": pct}),
            "for which is a Continuation of International application filed on March 14, 2023 "
            "(International Application No. PCT/CN2023/081289); and")

        cip = _special_chain({"filing_route": "Bypass CIP", "case_advance_info": pct,
                              "priorities": [cn("202311832487.5", "2023-12-27"),
                                             cn("202210688095.5", "2022-06-17"),
                                             cn("202210689642.1", "2022-06-17")]})
        self.assertEqual(cip,
            "for which is a Continuation-in-part of International Application filed on March 14, 2023 "
            "(International Application No. PCT/CN2023/081289), claiming priority to Chinese "
            "Patent Application No. 202311832487.5 filed on December 27, 2023, Chinese Patent "
            "Application No. 202210688095.5 filed on June 17, 2022, and Chinese Patent "
            "Application No. 202210689642.1 filed on June 17, 2022; and")

        paris = _special_chain({"filing_route": "Paris", "priorities": [cn("202511106870.1", "2025-08-07")]})
        self.assertEqual(paris, "claiming priority to Chinese Patent Application No. 202511106870.1 "
                         "filed on August 7, 2025; and")

        stage = _special_chain({"filing_route": "371", "case_advance_info": pct,
                                "priorities": [priority, cn("202422933118.1", "2024-11-28")]})
        self.assertEqual(stage,
            "for which is a 371 national-stage entry of International Patent Application No. "
            "PCT/CN2023/081289, filed on March 14, 2023; which claims priority to Chinese "
            "Patent Application No. 202211097238.1 filed on September 8, 2022, and claims priority "
            "to Chinese Patent Application No. 202422933118.1 filed on November 28, 2024; and")
        self.assertEqual(_special_chain({"filing_route": "371", "case_advance_info": pct}),
            "for which is a 371 national-stage entry of International Patent Application No. "
            "PCT/CN2023/081289, filed on March 14, 2023; and")

    def test_special_templates_keep_customer_wording_and_signature_layout(self):
        for template in SPECIAL_ASSIGNMENT_TEMPLATES.values():
            with self.subTest(template=template.name):
                doc = Document(template)
                self.assertIn("{{APPLICATION_CHAIN}}", "\n".join(p.text for p in doc.paragraphs))
                self.assertNotIn("Application Number", "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells))
        company = Document(SPECIAL_ASSIGNMENT_TEMPLATES["公司转公司"])
        self.assertEqual(company.tables[0].rows[0].cells[0].text, "Assignor")
        self.assertEqual(company.tables[0].rows[0].cells[-1].text, "{{ASSIGNOR}}")
        inventor = Document(SPECIAL_ASSIGNMENT_TEMPLATES["个人转公司"])
        general = Document(ASSIGNMENT_TEMPLATE)
        self.assertEqual([t._tbl.xml for t in inventor.tables],
                         [t._tbl.xml for t in general.tables[1:3]])
        self.assertFalse(any(br.get(qn("w:type")) == "page"
                             for p in inventor.paragraphs
                             for br in p._p.findall(".//" + qn("w:br"))))

    def test_other_clients_still_use_general_template(self):
        fields = {
            "case_volume": "20752-0001US00", "application_type": "发明",
            "case_info": {"case_name_en": "A DIFFERENT DEVICE"},
            "applicants": [{"applicant_name_en": "OTHER CORP.", "address1": "1 Main Street"}],
            "inventors": [{"last_name_en": "DOE", "first_name_en": "Jane"}],
        }
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                  "assignee": "OTHER CORP."}, folder)
            _, text = doc_text(result["file_path"])
            self.assertIn("A DIFFERENT DEVICE", text)
            self.assertNotIn("{{", text)

    def test_20618_uses_both_templates_without_sample_case_data(self):
        fields = {
            "case_volume": "20618-9999US00",
            "application_type": "发明",
            "case_info": {"case_name_en": "A NEW MEDICAL DEVICE"},
            "applicants": [{"applicant_name_en": "NEW MEDICAL CORP.",
                            "address1": "100 Main Street, Building 123, Industrial Development Zone",
                            "address2": "45 Riverside Avenue, Suite 321", "city_en": "Boston",
                            "province_city_en_us": "Massachusetts", "postcode": "02101",
                            "country_en_us": "US"}],
            "inventors": [{"last_name_en": name, "first_name_en": "Test",
                           "employer": "NEW MEDICAL CORP.",
                           "residence_en": ("1 Very Long Main Street, Building 123, Industrial Development Zone, "
                                            "45 Riverside Avenue, Suite 321, Boston, Massachusetts 02101, US")
                           if idx == 1 else f"{idx} Test Street, Boston, US"}
                          for idx, name in enumerate(("SMITH", "LEE", "KIM"), 1)],
            "assignments": [{"assignment_type": "公司转公司",
                             "assigning_company": "AMERICA HOLDINGS INC.",
                             "company_representative": "ALEX WU",
                             "representative_title": "CEO",
                             "company_address1": "100 Main Street",
                             "company_city_en": "Boston",
                             "company_country_en": "US",
                             "assignee": "NEW MEDICAL CORP."}],
            "parent_applications": [{"Application No": "18/123,456",
                                     "Prior Application Number": "18/123,456",
                                     "Continuity Type": "Continuation of",
                                     "Filing Date": "2024-01-01"}],
            "case_advance_info": {"pct_app_no": "PCT/CN2023/987654",
                                  "pct_app_date": "2023-01-01"},
            "priorities": [{"priority_no": "CN 202511280001.X", "priority_country_id": "CN",
                            "priority_date": "2025-11-28"}],
        }
        with tempfile.TemporaryDirectory() as folder:
            with patch("submission.assignment_generator.Document", wraps=Document) as load_template:
                files = generate_assignments(fields, folder)
            self.assertEqual([Path(c.args[0]).name for c in load_template.call_args_list],
                             ["Assignment-20618-inventors.docx", "Assignment-20618-company.docx"])
            self.assertEqual(len(files), 2)
            inventor_doc, inventor_text = doc_text(files[0]["file_path"])
            company_doc, company_text = doc_text(files[1]["file_path"])
            self.assertEqual(len(inventor_doc.tables), 3)
            self.assertEqual(len(company_doc.tables), 4)
            self.assertFalse(any(br.get(qn("w:type")) == "page"
                                 for p in inventor_doc.paragraphs
                                 for br in p._p.findall(".//" + qn("w:br"))))
            for table in inventor_doc.tables:
                first_p = table.rows[0].cells[0].paragraphs[0]
                self.assertIsNotNone(first_p._p.find(".//" + qn("w:pageBreakBefore")))
            self.assertGreater(len(inventor_doc.tables[0].rows), 5)  # 长地址续行
            self.assertEqual(len(inventor_doc.tables[0].columns), 3)
            self.assertIn("1 Very Long Main Street", inventor_doc.tables[0].rows[3].cells[-1].text)
            self.assertIn("2 Test Street", "\n".join(c.text for r in inventor_doc.tables[1].rows for c in r.cells))
            self.assertEqual(company_doc.tables[0].rows[0].cells[-1].text, "AMERICA HOLDINGS INC.")
            self.assertIn("for which is a continuation of", inventor_text)
            self.assertIn("for which is a continuation of", company_text)
            for text in (inventor_text, company_text):
                self.assertIn("20618-9999US00", text)
                self.assertIn("A NEW MEDICAL DEVICE", text)
                self.assertIn("U.S. Application No. 18/123,456", text)
                self.assertIn("PCT/CN2023/987654", text)
                self.assertIn("filed on January 1, 2024", text)
                self.assertIn("filed on January 1, 2023", text)
                self.assertIn("Chinese Patent Application No. 202511280001.X filed on November 28, 2025", text)
                self.assertNotIn("2025-11-28", text)
                self.assertNotIn("{{", text)
                for old in ("20618-0168US04", "SHANGHAI UNITED", "UNITED IMAGING",
                            "PCT/CN2016/085434", "No.2258 Chengbei"):
                    self.assertNotIn(old, text)
            self.assertIn("SMITH, Test", inventor_text)
            self.assertIn("ALEX WU", company_text)

            single = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                  "assignor_inventors": ["SMITH, Test"],
                                                  "assignee": "NEW MEDICAL CORP."}, Path(folder) / "single")
            single_doc, _ = doc_text(single["file_path"])
            self.assertEqual(len(single_doc.tables), 1)


if __name__ == "__main__":
    unittest.main()
