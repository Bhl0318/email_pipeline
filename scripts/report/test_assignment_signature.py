"""通用 Assignment 签字区回归测试；不连接数据库。"""

import tempfile
import unittest
from pathlib import Path

from docx import Document

from submission.assignment_generator import generate_assignment, generate_assignments
from submission.common import resolve_assignee, address_lines


class GeneralAssignmentSignatureTests(unittest.TestCase):
    def test_missing_nonapplicant_employer_address_blocks_without_deleting_old_file(self):
        fields = {
            "case_volume": "99999-0008US00", "application_type": "发明",
            "case_info": {"case_name_en": "TEST DEVICE"},
            "applicants": [{"applicant_name_en": "APPLICANT A", "address1": "A Road"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                           "employer": "COMPANY X", "address1": "Old Inventor Address"}],
            "employer_info": {"COMPANY X": {"applicant_name_en": "COMPANY X"}},
            "assignments": [{"assignment_type": "公司转公司",
                             "assigning_company": "COMPANY X", "company_address1": "X Road",
                             "assignee": "APPLICANT A"}],
        }
        # 公司转公司段有地址时可作为 X 的来源；清空后不能借用 A 或发明人旧住址。
        fields["assignments"][0]["company_address1"] = ""
        with tempfile.TemporaryDirectory() as folder:
            old = Path(folder) / "old-Assignment.docx"
            old.write_bytes(b"old")
            with self.assertRaisesRegex(ValueError, "COMPANY X.*缺少英文住址"):
                generate_assignments(fields, folder)
            self.assertEqual(old.read_bytes(), b"old")

    def test_missing_company_assignor_address_blocks(self):
        fields = {"case_volume": "99999-0009US00", "application_type": "发明",
                  "case_info": {"case_name_en": "TEST DEVICE"},
                  "applicants": [{"applicant_name_en": "APPLICANT A", "address1": "A Road"}],
                  "inventors": [{"last_name_en": "LI", "first_name_en": "Alice"}]}
        assignment = {"assignment_type": "公司转公司", "assigning_company": "COMPANY X",
                      "assignee": "APPLICANT A"}
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "COMPANY X.*缺少英文地址"):
                generate_assignment(fields, assignment, folder)

    def test_unknown_assignee_does_not_inherit_first_applicant_address(self):
        fields = {"case_volume": "99999-0010US00", "application_type": "发明",
                  "case_info": {"case_name_en": "TEST DEVICE"},
                  "applicants": [{"applicant_name_en": "APPLICANT A", "address1": "A Road"}],
                  "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                                 "employer": "APPLICANT A"}]}
        assignment = {"assignment_type": "个人转公司", "assignee": "COMPANY Y",
                      "assignor_inventors": ["LI, Alice"]}
        with tempfile.TemporaryDirectory() as folder:
            with self.assertRaisesRegex(ValueError, "COMPANY Y.*没有可匹配的地址"):
                generate_assignment(fields, assignment, folder)

    def test_inventor_employed_by_applicant_uses_employer_panel_address(self):
        fields = {
            "case_volume": "99999-0004US00", "application_type": "发明",
            "case_info": {"case_name_en": "TEST DEVICE"},
            "applicants": [{"applicant_name_en": "APPLICANT CO., LTD.",
                            "address1": "Company Street", "city_en": "Shanghai"}],
            "inventors": [
                {"last_name_en": "LI", "first_name_en": "Alice",
                 "employer": "Applicant Co Ltd", "residence_en": "Alice Residence, Suzhou, China"},
                {"last_name_en": "WU", "first_name_en": "Bob",
                 "employer": "APPLICANT CO., LTD.",
                 "address1": "Bob Home Road", "address2": "Nanjing, China"},
            ],
            "employer_info": {"APPLICANT CO., LTD.": {
                "residence_en": "Panel Office Road, Shanghai, China"}},
        }
        assignment = {"assignment_type": "个人转公司", "assignee": "APPLICANT CO., LTD.",
                      "assignor_inventors": ["LI, Alice", "WU, Bob"]}
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, assignment, folder)
            doc = Document(result["file_path"])
            for name in ("LI, Alice", "WU, Bob"):
                table = next(t for t in doc.tables
                             if any(name in cell.text for row in t.rows for cell in row.cells))
                signature = " ".join(cell.text for row in table.rows for cell in row.cells)
                self.assertIn("Panel Office Road", signature)
                self.assertNotIn("Company Street", signature)
                self.assertNotIn("Alice Residence", signature)
                self.assertNotIn("Bob Home Road", signature)
            self.assertFalse(result["warnings"])

    def test_non_applicant_employer_uses_employer_panel_address(self):
        fields = {
            "case_volume": "99999-0005US00", "application_type": "发明",
            "case_info": {"case_name_en": "TEST DEVICE"},
            "applicants": [{"applicant_name_en": "APPLICANT CO.", "address1": "Applicant Road"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                           "employer": "EMPLOYER CO.", "residence_en": "Alice Home"}],
            "assignments": [{"assignment_type": "公司转公司",
                             "assigning_company": "EMPLOYER CO.",
                             "company_address1": "Employer Office"}],
            "employer_info": {"EMPLOYER CO.": {"residence_en": "Panel Employer HQ"}},
        }
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                  "assignee": "EMPLOYER CO.",
                                                  "assignor_inventors": ["LI, Alice"]}, folder)
            doc = Document(result["file_path"])
            signature = " ".join(cell.text for table in doc.tables
                                 for row in table.rows for cell in row.cells)
            self.assertIn("Panel Employer HQ", signature)
            self.assertNotIn("Employer Office", signature)
            self.assertNotIn("Alice Home", signature)

    def test_applicant_employer_address_does_not_need_inventor_address(self):
        for case_no in ("99999-0006US00", "20618-0006US00"):
            with self.subTest(case_no=case_no), tempfile.TemporaryDirectory() as folder:
                fields = {
                    "case_volume": case_no, "application_type": "发明",
                    "case_info": {"case_name_en": "TEST DEVICE"},
                    "applicants": [{"applicant_name_en": "APPLICANT CO.",
                                    "address1": "Company Street"}],
                    "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                                   "employer": "APPLICANT CO."}],
                    "employer_info": {"APPLICANT CO.": {"residence_en": "Company Street"}},
                }
                result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                      "assignee": "APPLICANT CO.",
                                                      "assignor_inventors": ["LI, Alice"]}, folder)
                doc = Document(result["file_path"])
                signature = next(t for t in doc.tables
                                 if any("LI, Alice" in cell.text for row in t.rows for cell in row.cells))
                text = " ".join(cell.text for row in signature.rows for cell in row.cells)
                self.assertIn("Company Street", text)
                self.assertFalse(result["warnings"])

    def test_special_applicant_employer_uses_employer_panel_before_inventor_address(self):
        fields = {
            "case_volume": "20618-0007US00", "application_type": "发明",
            "case_info": {"case_name_en": "TEST DEVICE"},
            "applicants": [{"applicant_name_en": "APPLICANT CO.", "address1": "Company Street"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                           "employer": "APPLICANT CO.", "address1": "Alice Home Road"}],
            "employer_info": {"APPLICANT CO.": {"residence_en": "Company Street"}},
        }
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                  "assignee": "APPLICANT CO.",
                                                  "assignor_inventors": ["LI, Alice"]}, folder)
            doc = Document(result["file_path"])
            signature = next(t for t in doc.tables
                             if any("LI, Alice" in cell.text for row in t.rows for cell in row.cells))
            text = " ".join(cell.text for row in signature.rows for cell in row.cells)
            self.assertIn("Company Street", text)
            self.assertNotIn("Alice Home Road", text)

    def test_mailing_address_lines_have_no_trailing_commas(self):
        residence = "No. 601 Huangpu Avenue West, Tianhe District, Guangzhou, Guangdong, China"
        for case_no, width in (("99999-0003US00", 50), ("20618-0003US00", 36)):
            with self.subTest(case_no=case_no), tempfile.TemporaryDirectory() as folder:
                fields = {
                    "case_volume": case_no, "application_type": "发明",
                    "case_info": {"case_name_en": "TEST DEVICE"},
                    "applicants": [{"applicant_name_en": "EMPLOYER CO.", "address1": residence}],
                    "inventors": [{"last_name_en": "LI", "first_name_en": "Alice"}],
                }
                result = generate_assignment(fields, {"assignment_type": "个人转公司",
                                                      "assignee": "EMPLOYER CO."}, folder)
                doc = Document(result["file_path"])
                cells = [cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells]
                for line in address_lines(residence, "", width=width):
                    self.assertIn(line.rstrip(","), cells)
                    if line.endswith(","):
                        self.assertNotIn(line, cells)

    def test_inventor_to_employer_uses_independent_full_address(self):
        fields = {
            "case_volume": "99999-0002US00", "application_type": "发明",
            "case_info": {"case_name_en": "Test invention"},
            "applicants": [{"applicant_name_en": "APPLICANT CO.", "address1": "Applicant Road"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice",
                           "employer": "EMPLOYER CO., LTD."}],
            "employer_info": {"EMPLOYER CO., LTD": {
                "residence_en": "Full Employer Headquarters, No. 18 New Road, Shanghai, China"}},
            "assignments": [{"assignment_type": "公司转公司",
                             "assigning_company": "EMPLOYER CO., LTD.",
                             "company_address1": "Old Structured Address",
                             "assignee": "APPLICANT CO."}],
        }
        assignment = {"assignment_type": "个人转公司", "assignee": "EMPLOYER CO., LTD.",
                      "assignor_inventors": ["LI, Alice"]}
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, assignment, folder)
            doc = Document(result["file_path"])
            signature_text = " ".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            self.assertIn("Full Employer Headquarters", signature_text)
            self.assertIn("No. 18 New Road", signature_text)
            self.assertNotIn("Old Structured Address", signature_text)
            self.assertEqual(doc.tables[1].rows[2].cells[-1].text.strip(), "")  # Execution Date
        # 独立完整地址留空时，继续从公司转公司段的地址字段取值。
        fields["employer_info"]["EMPLOYER CO., LTD"]["residence_en"] = ""
        fallback = resolve_assignee(fields, "EMPLOYER CO., LTD.", prefer_employer_residence=True)
        self.assertIn("Old Structured Address", fallback["address_lines"])

    def test_company_assignor_once_and_execution_date_blank(self):
        fields = {
            "case_volume": "99999-0001US00",
            "application_type": "发明",
            "case_info": {"case_name_en": "Test invention"},
            "applicants": [{"applicant_name_en": "ASSIGNEE CO.",
                            "address1": "1 Main Street", "city_en": "Suzhou",
                            "country_en_us": "China"}],
            "inventors": [{"last_name_en": "LI", "first_name_en": "Alice"}],
        }
        assignment = {
            "assignment_type": "公司转公司",
            "assigning_company": "CRYSTAL PHARMACEUTICAL (SUZHOU) CO., LTD.",
            "company_representative": "bi",
            "company_address1": "No.1002, Dagunan Road",
            "company_city_en": "Tianjin",
        }
        with tempfile.TemporaryDirectory() as folder:
            result = generate_assignment(fields, assignment, folder)
            doc = Document(result["file_path"])
            tables = [table for table in doc.tables
                      if table.rows and table.rows[0].cells[0].text.strip() == "Assignor"]
            self.assertEqual(len(tables), 1)
            cells = [cell.text for row in tables[0].rows for cell in row.cells]
            self.assertIn("CRYSTAL PHARMACEUTICAL (SUZHOU) CO., LTD.", cells)
            self.assertIn("bi", cells)
            date_index = cells.index("Execution Date:")
            self.assertEqual(cells[date_index + 1], "")
            self.assertNotIn("转让公司代表签字日期", "".join(cells))


if __name__ == "__main__":
    unittest.main()
