#!/usr/bin/env python3
"""把结项 V2 Word 转成供人工校对的 Markdown 草稿；不覆盖线上模板。"""

import json
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "outputs" / "结项V2规范化草表.json"
WORD_DIR = ROOT / "最新版本" / "模板"
OUT_DIR = ROOT / "outputs" / "结项V2模板草稿"
FIELDS = {
    "case_volume": "代理所案号",
    "case_volume_customer": "客户案号",
    "customer_name": "客户名称",
    "case_name": "申请名称",
    "案件名称": "申请名称",
    "app_no": "申请号",
    "申请号": "申请号",
    "递交日": "递交日期",
    "申请日": "申请日",
    "处理事项完成日": "完成日期",
    "doc_date": "官文日期",
    "app_dateA12M": "临时申请到期日",
}
PLACEHOLDER = re.compile(r"《([^《》]+)》")
DATE = re.compile(r"\b20\d{2}\s*[年./-]\s*\d{1,2}\s*[月./-]\s*\d{1,2}\s*日?")


def _normal(text, unknown):
    text = text.strip().replace("\u00a0", " ")

    def repl(m):
        name = m.group(1).strip()
        if name not in FIELDS:
            unknown.add(name)
            return m.group(0)
        return "{{" + FIELDS[name] + "}}"

    text = PLACEHOLDER.sub(repl, text)
    text = text.replace("七星天", "{{签名品牌}}")
    return text


def _as_md(path):
    document = Document(path)
    blocks, unknown = [], set()
    for element in document.element.body:
        tag = element.tag.rsplit("}", 1)[-1]
        if tag == "p":
            s = _normal(Paragraph(element, document).text, unknown)
            if not s:
                continue
            if not blocks and (DATE.fullmatch(s) or s == "电子邮件送达"):
                continue
            if s == "电子邮件送达" and len(blocks) < 3:
                continue
            if s == "{{客户名称}}" and len(blocks) < 3:
                continue
            if s.startswith("Privileged and Confidential") or s.startswith("机密并受律师客户"):
                # 使用和现有 OA/结项 MD 一致的标准保密标头。
                continue
            if s.startswith("关于："):
                s = "**" + s + "**"
            blocks.append(s)
        elif tag == "tbl":
            table = Table(element, document)
            rows = []
            for row in table.rows:
                cells = [_normal(c.text, unknown).replace("\n", "<br>").replace("|", "\\|")
                         for c in row.cells]
                while cells and not cells[-1]:
                    cells.pop()
                if cells:
                    rows.append(cells)
            width = max((len(row) for row in rows), default=0)
            if width:
                table_lines = ["| " + " | ".join("" for _ in range(width)) + " |",
                               "| " + " | ".join("---" for _ in range(width)) + " |"]
                for row in rows:
                    table_lines.append("| " + " | ".join(row + [""] * (width - len(row))) + " |")
                blocks.append("\n".join(table_lines))
    result = "***机密并受律师客户通讯特权保护***\n***Privileged and Confidential; Attorney-Client Communication***\n\n"
    result += "\n\n".join(blocks) + "\n"
    return result, sorted(unknown), sorted(set(DATE.findall(result)))


def main():
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    report = []
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    seen = set()
    for r in data["records"]:
        region = r["sheet"].rsplit("-", 1)[-1]
        old = r["template_md"]
        if region not in ("US", "PCT") or not old or old in seen:
            continue
        seen.add(old)
        tags = "".join(re.findall(r"【([^】]+)】", r["template_label"])).replace(" ", "")
        candidates = [p for p in WORD_DIR.rglob("*.docx")
                      if region in p.parent.name
                      and tags == "".join(re.findall(r"【([^】]+)】", p.name)).replace(" ", "")]
        if len(candidates) != 1:
            raise ValueError(f"Word 映射不唯一：{r['sheet']}:{r['excel_row']} {tags!r} -> {candidates}")
        body, unknown, dates = _as_md(candidates[0])
        name = old[:-3] + "_V2.md"
        (OUT_DIR / name).write_text(body, encoding="utf-8")
        report.append({"template": name, "word": str(candidates[0].relative_to(ROOT)),
                       "unknown_placeholders": unknown, "fixed_dates_in_body": dates})
    generic_word = WORD_DIR / "待填模板-仅表头版-示例.docx"
    generic_body, generic_unknown, generic_dates = _as_md(generic_word)
    (OUT_DIR / "结项_待填模板_V2.md").write_text(generic_body, encoding="utf-8")
    report.append({"template": "结项_待填模板_V2.md",
                   "word": str(generic_word.relative_to(ROOT)),
                   "unknown_placeholders": generic_unknown,
                   "fixed_dates_in_body": generic_dates,
                   "requires_manual_body": True})
    (OUT_DIR / "转换校对清单.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"转换 {len(report) - 1} 个 US/PCT Word + 1 个待填模板草稿：{OUT_DIR}")
    print(f"含固定日期 {sum(bool(x['fixed_dates_in_body']) for x in report)}；"
          f"含未映射占位符 {sum(bool(x['unknown_placeholders']) for x in report)}")


if __name__ == "__main__":
    main()
