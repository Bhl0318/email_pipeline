#!/usr/bin/env python3
"""生成文档不出现字面 'None' 的回归测试（null 字段应渲染成空白，不是 'None'）。

  cd scripts/report && python3 test_no_none_leak.py
"""
import os
import sys
import tempfile

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, SCRIPT_DIR)

from submission.common import blank_none
from submission.fields_mapping import fields_to_case_info
from submission.confirmation_sheet import generate_confirmation_sheet
from docx import Document

_passed = 0


def ok(cond, msg):
    global _passed
    assert cond, f"FAIL: {msg}"
    _passed += 1
    print(f"  ok - {msg}")


def _has_none(o):
    if o is None:
        return True
    if isinstance(o, dict):
        return any(_has_none(v) for v in o.values())
    if isinstance(o, list):
        return any(_has_none(v) for v in o)
    return False


FIELDS = {
    "case_volume": "X-001US00", "application_type": "发明", "ae_methods": "无加速",
    "case_info": {"case_name_en": "TEST", "case_name_cn": None},
    "applicants": [{"applicant_name_en": "ACME CO., LTD.", "applicant_name_cn": None,
                    "address1": "Rd 1", "address2": None, "city_en": "Beijing",
                    "province_city_en_us": None, "postcode": "100000", "country_en_us": "China", "entity_en": "Large"}],
    "inventors": [{"last_name_en": "W", "first_name_en": "H", "full_name_en": "W, H", "name_cn": None, "employer": "ACME CO., LTD."}],
    "priorities": [], "parent_applications": [], "filing_route": "371",
    "case_advance_info": {"pct_app_no": "PCT/X", "pct_app_date": None, "pct_pub_no": None},
}


def test_blank_none():
    ok(blank_none({"a": None, "b": [None, {"c": None}]}) == {"a": "", "b": ["", {"c": ""}]},
       "blank_none 深度把 None 转成空串")


def test_case_info_no_none():
    ci = fields_to_case_info(FIELDS)
    ok(not _has_none(ci), "fields_to_case_info 输出无 None（ADS/POA/Transmittal 不泄漏 None）")


def test_confirmation_no_none_text():
    out = tempfile.mkdtemp(prefix="none_")
    r = generate_confirmation_sheet(FIELDS, out)
    d = Document(r["file_path"])
    cells = [c.text.strip() for t in d.tables for row in t.rows for c in row.cells]
    paras = [p.text for p in d.paragraphs]
    ok("None" not in cells, "确认函单元格无字面 'None'")
    ok(not any("None" in p for p in paras), "确认函正文无字面 'None'")


if __name__ == "__main__":
    test_blank_none()
    test_case_info_no_none()
    test_confirmation_no_none_text()
    print(f"\nAll {_passed} checks passed.")
