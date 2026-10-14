"""手工转换材料并上传外部前端分析接口的离线测试。"""

import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

sys.path.insert(0, os.path.dirname(__file__))
from qdfx_external_manual import submit_markdown_files


class ManualSubmitTests(unittest.TestCase):
    @patch("qdfx_external_manual._connection_config", return_value=("test-token", "http://example.test/assist"))
    def test_uploads_original_markdown_files_and_returns_parse_id(self, config):
        with tempfile.TemporaryDirectory() as tmp:
            paths = [Path(tmp) / name for name in ("oa.md", "application.md", "ref1.md", "ref2.md")]
            for path in paths:
                path.write_text(f"# {path.stem}\n原始内容", encoding="utf-8")

            def post(url, *, files, data, headers, timeout):
                self.assertEqual(url, "http://example.test/assist/qdfx/external/parse")
                self.assertEqual([field for field, _ in files], ["scyj", "sqwj", "dbwj", "dbwj"])
                self.assertEqual([entry[1].read().decode("utf-8") for _, entry in files],
                                 [path.read_text(encoding="utf-8") for path in paths])
                self.assertEqual(data, {"caseNo": "90009-0001CN00", "createdBy": "writeSYS"})
                self.assertEqual(headers, {"X-Api-Token": "test-token"})
                return Mock(status_code=202, json=lambda: {"code": 202, "data": {"parseId": "p1"}})

            result = submit_markdown_files("90009-0001CN00", *paths[:2], paths[2:], post=post,
                                           output_dir=Path(tmp) / "converted")
        self.assertEqual(result["parseId"], "p1")

    @patch("qdfx_external_manual._connection_config", return_value=("test-token", "http://example.test/assist"))
    @patch("mistral_ocr.pdf_to_markdown", return_value="# 审查意见\n\nOCR 正文")
    def test_converts_pdf_docx_and_txt_before_upload(self, ocr, config):
        from docx import Document

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            oa = root / "审查意见.pdf"
            oa.write_bytes(b"PDF input for mocked OCR")
            app = root / "申请文件.docx"
            doc = Document()
            doc.add_heading("申请文件", level=1)
            doc.add_paragraph("说明书正文")
            table = doc.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "序号"
            table.cell(0, 1).text = "内容"
            table.cell(1, 0).text = "1"
            table.cell(1, 1).text = "技术方案"
            doc.save(app)
            ref1 = root / "对比文件1.txt"
            ref1.write_text("对比文件正文", encoding="utf-8")
            ref2 = root / "对比文件2.md"
            ref2.write_text("# 已有 Markdown", encoding="utf-8")

            def post(url, *, files, data, headers, timeout):
                self.assertEqual([field for field, _ in files], ["scyj", "sqwj", "dbwj", "dbwj"])
                uploaded = [entry[1].read().decode("utf-8") for _, entry in files]
                self.assertEqual(uploaded[0], "# 审查意见\n\nOCR 正文")
                self.assertIn("# 申请文件\n\n说明书正文", uploaded[1])
                self.assertIn("| 序号 | 内容 |", uploaded[1])
                self.assertEqual(uploaded[2:], ["对比文件正文", "# 已有 Markdown"])
                self.assertTrue(all(name.endswith(".md") for _, (name, _, _) in files))
                return Mock(status_code=202, json=lambda: {"code": 202, "data": {"parseId": "p2"}})

            result = submit_markdown_files("90009-0001CN00", oa, app, [ref1, ref2], post=post,
                                           output_dir=root / "converted")
            self.assertEqual(result["parseId"], "p2")
            self.assertEqual(len(list((root / "converted" / "90009-0001CN00").glob("*.md"))), 3)
            ocr.assert_called_once_with(str(oa))


if __name__ == "__main__":
    unittest.main()
